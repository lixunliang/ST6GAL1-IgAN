# -*- coding: utf-8 -*-
"""Build cis instruments for Olink proteins from Sun et al. 2023 (Nature) Supplementary Tables 9 and 16.

Input : raw/reference/ukbppp_sun2023/41586_2023_6592_MOESM3_ESM.xlsx
Output: extracted/olink_cis_instruments.tsv
  ST9  = lead cis-pQTLs (BETA w.r.t. A1, per allele)              -> source "ST9_lead"
  ST16 = SuSiE conditionally independent cis signals; effects are on a
         standardized-genotype scale and are converted to per-allele by dividing
         beta and SE by sqrt(2p(1-p))                                -> source "ST16_cond"
Only single-nucleotide variants are kept. MHC = chr6:25-34.5 Mb (GRCh37).
"""
import collections
import math
import os

import openpyxl

PROJ = os.environ.get("MR_PROJ", r"F:\孟德尔")
XLSX = os.path.join(PROJ, "raw", "reference", "ukbppp_sun2023", "41586_2023_6592_MOESM3_ESM.xlsx")
OUT = os.path.join(PROJ, "extracted", "olink_cis_instruments.tsv")


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    rows = list(wb["ST9"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r[0] and str(r[0]).startswith("Variant ID"))
    ix = {str(h): i for i, h in enumerate(rows[hi])}
    out = []
    for r in rows[hi + 1:]:
        if r[ix["cis/trans"]] != "cis":
            continue
        v = str(r[0]).split(":")
        pid = str(r[ix["UKBPPP ProteinID"]])
        out.append([pid, pid.split(":")[0], "ST9_lead", r[ix["rsID"]] or "", v[0], v[1], r[ix["GENPOS (hg38)"]], v[2], v[3],
                    float(r[ix["A1FREQ (discovery)"]]), float(r[ix["BETA (discovery, wrt. A1)"]]), float(r[ix["SE (discovery)"]]),
                    float(r[ix["log10(p) (discovery)"]]), r[ix["MHC"]]])
    mhc = {o[0]: o[13] for o in out}
    seen = {(o[0], o[4], o[5]) for o in out}
    n16 = 0
    for k, r in enumerate(wb["ST16"].iter_rows(values_only=True)):
        if k < 5 or not r[0] or r[9] != "cis":
            continue
        v = str(r[1]).split(":")
        pid = str(r[0])
        if len(v) < 4 or (pid, v[0], v[1]) in seen:
            continue
        try:
            b, se, lp, af = float(r[5]), float(r[6]), float(r[7]), float(r[8])
        except (TypeError, ValueError):
            continue
        s = math.sqrt(2 * af * (1 - af))
        if s <= 0:
            continue
        inm = 1 if (v[0] == "6" and 25_000_000 <= int(v[1]) <= 34_500_000) else 0
        out.append([pid, pid.split(":")[0], "ST16_cond", r[2] or "", v[0], v[1], "", v[2], v[3], af, b / s, se / s, lp, mhc.get(pid, inm)])
        seen.add((pid, v[0], v[1]))
        n16 += 1
    out = [o for o in out if len(str(o[7])) == 1 and len(str(o[8])) == 1]
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("ProteinID\tGene\tsource\trsid\tCHR\tPOS37\tPOS38\tOA\tEA\tEAF\tBeta\tSE\tlog10p\tMHC\n")
        for o in out:
            f.write("\t".join(f"{x:.6g}" if isinstance(x, float) else str(x) for x in o) + "\n")
    print(f"{len(out)} instruments for {len({o[0] for o in out})} proteins ({n16} from ST16)")
    print(collections.Counter(o[2] for o in out))


if __name__ == "__main__":
    main()
